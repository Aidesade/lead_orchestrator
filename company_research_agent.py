# -*- coding: utf-8 -*-
r"""
Двухстадийный агент-ресёрчер компаний на Claude Agent SDK, режим ПРЕСЕЙЛ.

  STAGE 1 (triage)   : грязный текст / ИНН  ->  чистый JSON {company_name, inn, aspects}
  STAGE 2 (report)   : handle компании  ->  ресёрч + ДВА пресейл-документа .docx:
        (1) КАРТА БИЗНЕС-ПРОЦЕССОВ — профиль → as-is процессы → боли →
            ИИ-решение → to-be, сводная таблица, дорожная карта, специфика 44/223-ФЗ;
        (2) КАРТА РОЛЕЙ И КОНТАКТОВ · ПРЕСЕЙЛ — оргструктура, ЛПР, профильные отделы
            под внедрение ИИ, филиалы, офиц. контакты, план захода.

Stage 2 — это агент, у которого есть инструменты:
  - mcp__research__deep_research        : НАСТОЯЩИЙ глубокий ресёрч (движок deep_research_engine)
        официальная база (ГИР БО + Dadata + Checko, как раньше) +
        РЕАЛЬНЫЙ краул сайта (филиалы+директора+телефоны, руководство, контакты,
        соцсети, проектный институт) + ЕИС/госзакупки по ИНН + суды/СМИ + hh.ru;
        supervisor с параллельными коллекторами и петля ЦЕЛЕВОГО добора под пустые
        ячейки таблиц. Каждая строка несёт source URL. Crawl4AI PRIMARY -> HTTP-фолбэк.
  - WebSearch / WebFetch (встроенные)   : ТОЧЕЧНАЯ доверка и закрытие остаточных пробелов
        КАЖДАЯ спорная ссылка/контакт проверяется WebFetch'ем (анти-галлюцинация)
  - mcp__research__save_process_map_docx   : КАРТА БИЗНЕС-ПРОЦЕССОВ -> .docx
  - mcp__research__save_roles_contacts_docx: КАРТА РОЛЕЙ И КОНТАКТОВ -> .docx

Компания, определённая на первом проходе (stage 1), ЯВНО передаётся во второй.

Запуск:
  py C:/.../company_research_agent.py "АО Рязаньавтодор ИНН 6234065445"
  py .../company_research_agent.py            # без аргумента — интерактивный режим
  py .../company_research_agent.py --contacts "<...>"   # старый режим: краткий отчёт по контактам

Опц. через env: PRESALE_VENDOR (строка «Подготовлено для», по умолчанию пусто),
PRESALE_PLATFORM_DESC (описание платформы в предмете отчёта).
Ключи (опционально — ГИР БО работает и без них):  setx DADATA_TOKEN <...>   setx CHECKO_TOKEN <...>
Зависимости:  pip install claude-agent-sdk python-docx
"""
import json
import os
import re
import sys
import warnings

# Глушим косметический RequestsDependencyWarning (chardet 7.x вне диапазона requests) ДО импорта requests.
warnings.filterwarnings("ignore", message=r".*doesn't match a supported version.*")

import anyio
from claude_agent_sdk import (
    query,                 # async one-shot -> AsyncIterator[Message]
    tool,                  # @tool — in-process MCP-инструмент
    create_sdk_mcp_server, # in-process MCP-сервер (без subprocess)
    ClaudeAgentOptions,    # опции (поля snake_case)
    ClaudeSDKClient,       # многоходовый клиент (context manager)
    AssistantMessage, TextBlock, ToolUseBlock, ResultMessage,
)

# консоль Windows = cp1251 и роняет вывод на ₽/кириллице -> принудительно utf-8
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# Все модули скилла лежат в этом же каталоге — добавляем его в sys.path.
SCRIPTS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPTS)

from revenue_enrich import girbo_revenue          # noqa: E402  ГИР БО — бесплатно, без ключа

# Dadata/Checko опциональны: импортируем мягко, чтобы файл работал даже без них.
try:                                              # noqa: E402
    from dadata_enrich import DadataClient, extract_lpr
except Exception:
    DadataClient = None
try:                                              # noqa: E402
    from checko_enrich import CheckoClient, extract_company_contacts
except Exception:
    CheckoClient = None

# Модели. Алиасы 'opus'/'sonnet' резолвятся в текущие дефолты аккаунта.
MODEL = "opus"          # stage 2 — пресейл-документы (ресёрч + синтез)
FAST_MODEL = "sonnet"   # stage 1 — дешёвый триаж

DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")

# --- Параметры пресейла (переопределяются переменными окружения) ------------
VENDOR = os.environ.get("PRESALE_VENDOR", "")   # пусто -> строка «Подготовлено для» не выводится
PLATFORM_DESC = os.environ.get(
    "PRESALE_PLATFORM_DESC",
    "корпоративная on-premises LLM-платформа (RAG-база знаний, автономные ИИ-агенты, "
    "ИИ-Коуч «Наставник»)",
)
SOURCES_NOTE_DEFAULT = (
    "только публично доступные данные (оф. сайт, ЕГРЮЛ/Rusprofile, zakupki.gov.ru, "
    "реестр контрактов 44-ФЗ/223-ФЗ, судебная практика, отраслевая нормативка, вакансии)"
)
ETHICS_DEFAULT = (
    "персональные и семейные данные не собирались; вход — строго через официальные "
    "каналы и конкурентную процедуру (44-ФЗ/223-ФЗ), без обхода закупок через личные связи"
)
PURPOSE_DEFAULT = (
    "оргструктура и официальные деловые контакты для последующей проработки. "
    "Только публичные источники."
)


def _digits(s):
    return re.sub(r"\D", "", str(s or ""))


# ===========================================================================
# ОФИЦИАЛЬНАЯ БАЗА ПО ИНН (синхронный сбор из модулей скилла; блокирующий)
# ===========================================================================
def research_company(company_name: str, inn: str = "", aspects: str = "") -> dict:
    """ИНН/название -> сводный payload по модулям скилла. Блокирующий (сеть)."""
    inn = _digits(inn)
    card: dict = {}        # карточка Dadata (ЕГРЮЛ/ЕГРИП)
    contacts: dict = {}    # контакты Checko
    notes: list = []

    dada = None
    if DadataClient and os.environ.get("DADATA_TOKEN"):
        try:
            dada = DadataClient()
        except Exception as e:
            notes.append(f"Dadata недоступен: {e}")
    elif not os.environ.get("DADATA_TOKEN"):
        notes.append("DADATA_TOKEN не задан — карточка ЕГРЮЛ/ЛПР пропущена.")

    # ИНН не дан -> резолвим по названию через Dadata suggest
    if not inn and dada and company_name:
        try:
            cands = dada.suggest(company_name, count=5)
            if cands:
                inn = _digits((cands[0] or {}).get("inn"))
                if len(cands) > 1:
                    notes.append("Несколько кандидатов по названию — взят первый; уточни ИНН.")
            else:
                notes.append("Dadata не нашёл компанию по названию.")
        except Exception as e:
            notes.append(f"Резолв ИНН по названию не удался: {e}")

    if not inn:
        notes.append("ИНН не определён — выручка/контакты по ИНН недоступны.")

    # 1) карточка ЕГРЮЛ/ЕГРИП (ЛПР, ОКВЭД, адрес, статус) — Dadata findById
    if inn and dada:
        try:
            data = dada.find_by_inn(inn)
            card = extract_lpr(data) if data else {}
            if not card:
                notes.append("ИНН не найден в ЕГРЮЛ/ЕГРИП (Dadata).")
        except Exception as e:
            notes.append(f"Карточка Dadata не получена: {e}")

    # 2) ОФИЦИАЛЬНАЯ выручка (бухгалтерская, стр. 2110) — ГИР БО, бесплатно
    rev = girbo_revenue(inn) if inn else {}

    # 3) контакты (тел/email/сайт/ЛПР) — Checko /company
    if inn and CheckoClient and os.environ.get("CHECKO_TOKEN"):
        try:
            data = CheckoClient().company(inn)
            contacts = extract_company_contacts(data) if data else {}
        except Exception as e:
            notes.append(f"Контакты Checko не получены: {e}")
    elif inn and not os.environ.get("CHECKO_TOKEN"):
        notes.append("CHECKO_TOKEN не задан — контакты пропущены.")

    return {"company_name": company_name, "inn": inn, "aspects": aspects,
            "card": card, "revenue": rev, "contacts": contacts, "notes": notes}


def _render_official(p: dict) -> str:
    """payload официальной базы -> markdown (для подачи агенту/контактного режима)."""
    card, rev, c = p["card"], p["revenue"], p["contacts"]
    name = card.get("legal_name") or p["company_name"] or "Компания"
    out = [f"# {name}", ""]
    if p["inn"]:
        out.append(f"- ИНН {p['inn']}" + (f" · ОГРН {card['ogrn']}" if card.get("ogrn") else ""))
    if card.get("opf"):
        out.append(f"- Форма {card['opf']} · статус {card.get('state_status', '—')}")
    if card.get("okved"):
        out.append(f"- ОКВЭД {card['okved']} {card.get('okved_name', '')}".rstrip())
    if card.get("address"):
        out.append(f"- Адрес {card['address']}")
    if card.get("lpr_fio"):
        out.append(f"- Руководитель (ЛПР) {card['lpr_fio']} — {card.get('lpr_post', '')}".rstrip(" —"))

    out += ["", "## Выручка (бухгалтерская, стр. 2110)"]
    if rev.get("revenue") is not None:
        rub = f"{rev['revenue']:,}".replace(",", " ")
        src, srcname = rev.get("source_url") or "", rev.get("source_name", "ГИР БО ФНС")
        out.append(f"- {rub} ₽ за {rev.get('year', '—')} — источник: "
                   + (f"{srcname} {src}" if src else srcname))
        out.append("- ⚠️ Бухгалтерская выручка (начисление), НЕ оборот по расчётному счёту.")
    elif card.get("revenue") is not None:
        rub = f"{int(card['revenue']):,}".replace(",", " ")
        out.append(f"- ~{rub} ₽ (доходы по данным ФНС / Dadata) за {card.get('revenue_year', '—')}.")
    else:
        out.append("- Нет в открытом доступе (ИП без отчётности / банк / нет данных в ГИР БО).")

    if c:
        out += ["", "## Контакты (Checko / ЕГРЮЛ)"]
        if c.get("phone"):
            out.append(f"- Тел: {c['phone']}")
        if c.get("email"):
            kind = c.get("email_kind", "")
            out.append(f"- Email: {c['email']}" + (f" ({kind})" if kind else ""))
        if c.get("website"):
            out.append(f"- Сайт: {c['website']}")
        if c.get("ceo"):
            out.append(f"- ЛПР: {c['ceo']}")
    if p["notes"]:
        out += ["", "## Примечания"] + [f"- {n}" for n in p["notes"]]
    return "\n".join(out)


# ===========================================================================
# ИНСТРУМЕНТ 1 — НАСТОЯЩИЙ deep_research (движок deep_research_engine)
#   supervisor + параллельные коллекторы по источникам (офиц. база, сайт, ЕИС,
#   суды/СМИ, hh) + петля целевого добора под пустые ячейки таблиц + реальный
#   краул (Crawl4AI -> HTTP-фолбэк). Каждая строка несёт source URL. Сигнатура
#   (company_name, inn, aspects) СОХРАНЕНА — orchestrator/писатель не ломаются.
# ===========================================================================
@tool(
    "deep_research",
    "ГЛУБОКИЙ ресёрч компании по ИНН/названию: официальная база (ГИР БО/Dadata/Checko) "
    "+ РЕАЛЬНЫЙ сбор с сайта компании (филиалы и их директора+телефоны, руководство, "
    "контакты, соцсети), ЕИС/госзакупки по ИНН, суды/СМИ и hh.ru. Внутри — supervisor, "
    "параллельные коллекторы и петля ЦЕЛЕВОГО добора под незаполненные ячейки таблиц. "
    "Возвращает источникованный документ (markdown + структурный JSON; у каждой строки URL). "
    "Зови ОДИН раз в начале — он приносит сайт/филиалы/контакты, дальше WebFetch только для точечной доверки.",
    {
        "type": "object",
        "properties": {
            "company_name": {"type": "string", "description": "Название компании"},
            "inn": {"type": "string", "description": "ИНН 10/12 цифр, если известен"},
            "aspects": {"type": "string", "description": "Аспекты через запятую"},
        },
        "required": ["company_name"],
    },
)
async def deep_research(args):
    company_name = args["company_name"]
    inn = args.get("inn", "")
    aspects = args.get("aspects", "")
    # Движок — async (сам уводит блокирующее в потоки и капает конкуренцию).
    # Ленивый импорт рвёт цикл CRA<->движок; при ЛЮБОМ сбое — официальная база (пайплайн не падает).
    try:
        from deep_research_engine import deep_research as _engine
        text = await _engine(company_name, inn, aspects)
    except Exception as e:
        payload = await anyio.to_thread.run_sync(research_company, company_name, inn, aspects)
        text = (_render_official(payload)
                + f"\n\n[deep_research_engine недоступен: {str(e)[:120]} — вернулась только "
                  "официальная база; добери сайт/филиалы/контакты через WebSearch+WebFetch вручную.]")
    return {"content": [{"type": "text", "text": text}]}


# ===========================================================================
# РЕНДЕР .docx — общие хелперы python-docx
# ===========================================================================
def _safe_name(s: str) -> str:
    s = re.sub(r'[\\/:*?"<>|«»]', "", s or "Компания").strip()
    return re.sub(r"\s+", "_", s)[:90] or "Компания"


def _new_doc():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor

    doc = Document()
    sec = doc.sections[0]
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(1.5))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    pf = normal.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing = 1.05
    return doc, Pt, RGBColor


def _title_block(doc, Pt, RGBColor, title, subtitles, meta):
    """Заголовок отчёта: крупный титул + строки-подзаголовки + серый мета-блок."""
    GRAY = RGBColor(0x5A, 0x5A, 0x5A)
    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(2)
    tr = t.add_run(title)
    tr.bold = True
    tr.font.size = Pt(15)
    for s in subtitles:
        if not s:
            continue
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(1)
        sr = sp.add_run(s)
        sr.bold = True
        sr.font.size = Pt(10.5)
    for label, value in meta:
        if not value:
            continue
        mp = doc.add_paragraph()
        mp.paragraph_format.space_after = Pt(1)
        if label:
            lr = mp.add_run(f"{label}: ")
            lr.bold = True
            lr.font.size = Pt(8.5)
            lr.font.color.rgb = GRAY
        vr = mp.add_run(value)
        vr.font.size = Pt(8.5)
        vr.font.color.rgb = GRAY


def _h(doc, text, level=1):
    return doc.add_heading(text, level=level)


def _para(doc, Pt, text):
    p = doc.add_paragraph(text or "")
    p.paragraph_format.space_after = Pt(3)
    return p


def _bullet(doc, Pt, text):
    """Пункт списка; жирная лид-метка до первого ':' либо '—'."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    m = re.match(r"^(.*?[:—])(\s*)(.*)$", text or "", re.DOTALL)
    if m:
        p.add_run(m.group(1)).bold = True
        if m.group(3):
            p.add_run(" " + m.group(3))
    else:
        p.add_run(text or "")
    return p


def _labeled(doc, Pt, label, text):
    """Абзац вида «As-is: ...» — жирная метка, дальше обычный текст."""
    if not text:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"{label}: ").bold = True
    p.add_run(text)


def _table(doc, Pt, headers, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, htxt in enumerate(headers):
        cell = table.rows[0].cells[i]
        hp = cell.paragraphs[0]
        hp.paragraph_format.space_after = Pt(0)
        hr = hp.add_run(str(htxt))
        hr.bold = True
        hr.font.size = Pt(8.5)
    for row in rows:
        cells = table.add_row().cells
        for i in range(len(headers)):
            vp = cells[i].paragraphs[0]
            vp.paragraph_format.space_after = Pt(0)
            vr = vp.add_run(str(row[i] if i < len(row) else ""))
            vr.font.size = Pt(8.5)


def _save_doc(doc, path):
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc.save(path)
    return path


# ===========================================================================
# ДОКУМЕНТ 1 — КАРТА БИЗНЕС-ПРОЦЕССОВ
# ===========================================================================
PROCESS_MAP_SCHEMA = {
    "type": "object",
    "properties": {
        "org_name": {"type": "string", "description": "Полное наименование, напр. «АО «Рязаньавтодор»»"},
        "subject": {"type": "string",
                    "description": "Предмет: стратегия внедрения " + PLATFORM_DESC + " (можно уточнить под компанию)"},
        "sources_note": {"type": "string", "description": "Чем подтверждены данные (по умолчанию — публичные источники)"},
        "date": {"type": "string", "description": "Дата, напр. «июнь 2026 г.»"},
        "tldr": {"type": "array", "items": {"type": "string"},
                 "description": "Краткое резюме (TL;DR): 2–4 пункта — кто компания, где сильнее всего эффект ИИ, как продавать"},
        "key_findings": {"type": "array", "items": {"type": "string"},
                         "description": "Ключевые выводы: 3–6 пунктов с фактами и источниками"},
        "profile_table": {"type": "array", "items": {
            "type": "object",
            "properties": {"param": {"type": "string"}, "value": {"type": "string"}, "source": {"type": "string"}},
            "required": ["param", "value"]},
            "description": "Раздел «1. Профиль компании»: строки таблицы Параметр|Значение|Источник"},
        "financials": {"type": "array", "items": {"type": "string"},
                       "description": "Финансовые показатели (выручка/прибыль/госконтракты) с годом и источником"},
        "structure_qc": {"type": "array", "items": {"type": "string"},
                         "description": "Филиалы, структура, контроль качества (абзацы)"},
        "contracts_table": {"type": "array", "items": {
            "type": "object",
            "properties": {"year": {"type": "string"}, "subject": {"type": "string"},
                           "amount": {"type": "string"}, "source": {"type": "string"}},
            "required": ["subject"]},
            "description": "Ключевые госконтракты: строки Год|Предмет|Сумма|Источник"},
        "processes": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Напр. «2.1. Тендерная работа по 44-ФЗ»"},
                "as_is": {"type": "string", "description": "Как процесс устроен сейчас"},
                "pains": {"type": "string", "description": "Боли/узкие места (с цифрами/фактами, где есть)"},
                "ai_solution": {"type": "string",
                                "description": "ИИ-решение: RAG / автономный ИИ-агент / ИИ-Коуч «Наставник» — что делает"},
                "to_be": {"type": "string", "description": "To-be эффект (измеримо, где можно)"}},
            "required": ["title", "as_is", "ai_solution", "to_be"]},
            "description": "Раздел «2. Карта бизнес-процессов»: 6–9 процессов компании"},
        "summary_table": {"type": "array", "items": {
            "type": "object",
            "properties": {"process": {"type": "string"}, "pain": {"type": "string"},
                           "product": {"type": "string"}, "effect": {"type": "string"}},
            "required": ["process", "product", "effect"]},
            "description": "Раздел «3. Сводная таблица»: Процесс|Ключевая боль|Продукт/решение|Ожидаемый эффект"},
        "roadmap": {"type": "array", "items": {"type": "string"},
                    "description": "Раздел «4. Приоритизация и дорожная карта»: этапы 0→3 (демо→пилот→масштаб→зрелость)"},
        "gov_sale": {"type": "array", "items": {"type": "string"},
                     "description": "Раздел «5. Специфика продажи (44-ФЗ/223-ФЗ)»: путь продажи, ЛПР, on-prem УТП, НМЦК, триггеры"},
        "recommendations": {"type": "array", "items": {"type": "string"},
                            "description": "Рекомендации: с чего начать, на кого выходить, на чём делать акцент, метрики пилота"},
        "disclaimers": {"type": "array", "items": {"type": "string"},
                        "description": "Оговорки по достоверности данных (что не подтверждено, что оценочно)"},
        "filename": {"type": "string", "description": "Имя файла (опц.)"},
    },
    "required": ["org_name", "tldr", "processes"],
}


def _write_process_map_docx(payload: dict, path: str = None) -> str:
    """КАРТА БИЗНЕС-ПРОЦЕССОВ -> .docx. Возвращает путь к файлу."""
    doc, Pt, RGBColor = _new_doc()
    org = payload.get("org_name") or payload.get("company") or "Компания"

    _title_block(
        doc, Pt, RGBColor,
        "АНАЛИТИЧЕСКИЙ ОТЧЁТ",
        [f"Карта бизнес-процессов {org}",
         "as-is процессы → точки внедрения ИИ → to-be"],
        [("Подготовлено для", VENDOR),
         ("Предмет", payload.get("subject") or ("стратегия внедрения " + PLATFORM_DESC)),
         ("Источники", payload.get("sources_note") or SOURCES_NOTE_DEFAULT),
         ("Дата", payload.get("date") or "")],
    )

    if payload.get("tldr"):
        _h(doc, "Краткое резюме (TL;DR)", 1)
        for b in payload["tldr"]:
            _bullet(doc, Pt, b)

    if payload.get("key_findings"):
        _h(doc, "Ключевые выводы", 1)
        for b in payload["key_findings"]:
            _bullet(doc, Pt, b)

    _h(doc, "1. Профиль компании", 1)
    _table(doc, Pt, ["Параметр", "Значение", "Источник"],
           [[r.get("param", ""), r.get("value", ""), r.get("source", "")]
            for r in (payload.get("profile_table") or [])])
    if payload.get("financials"):
        _h(doc, "Финансовые показатели", 2)
        for b in payload["financials"]:
            _bullet(doc, Pt, b)
    if payload.get("structure_qc"):
        _h(doc, "Филиалы, структура и контроль качества", 2)
        for b in payload["structure_qc"]:
            _para(doc, Pt, b)
    if payload.get("contracts_table"):
        _h(doc, "Ключевые госконтракты", 2)
        _table(doc, Pt, ["Год", "Предмет", "Сумма", "Источник"],
               [[r.get("year", ""), r.get("subject", ""), r.get("amount", ""), r.get("source", "")]
                for r in payload["contracts_table"]])

    _h(doc, "2. Карта бизнес-процессов (as-is) → боли → ИИ-решение → to-be", 1)
    for pr in (payload.get("processes") or []):
        _h(doc, pr.get("title") or "Процесс", 3)
        _labeled(doc, Pt, "As-is", pr.get("as_is"))
        _labeled(doc, Pt, "Боли", pr.get("pains"))
        _labeled(doc, Pt, "ИИ-решение", pr.get("ai_solution"))
        _labeled(doc, Pt, "To-be эффект", pr.get("to_be"))

    if payload.get("summary_table"):
        _h(doc, "3. Сводная таблица: процесс → боль → продукт → эффект", 1)
        _table(doc, Pt, ["Процесс", "Ключевая боль", "Продукт / решение", "Ожидаемый эффект"],
               [[r.get("process", ""), r.get("pain", ""), r.get("product", ""), r.get("effect", "")]
                for r in payload["summary_table"]])

    if payload.get("roadmap"):
        _h(doc, "4. Приоритизация и дорожная карта внедрения", 1)
        for b in payload["roadmap"]:
            _para(doc, Pt, b)

    if payload.get("gov_sale"):
        _h(doc, "5. Специфика продажи госкомпании (44-ФЗ/223-ФЗ)", 1)
        for b in payload["gov_sale"]:
            _bullet(doc, Pt, b)

    if payload.get("recommendations"):
        _h(doc, "Рекомендации", 1)
        for b in payload["recommendations"]:
            _bullet(doc, Pt, b)

    if payload.get("disclaimers"):
        _h(doc, "Оговорки по достоверности данных", 1)
        for b in payload["disclaimers"]:
            _bullet(doc, Pt, b)

    if not path:
        os.makedirs(DOWNLOADS, exist_ok=True)
        path = os.path.join(DOWNLOADS,
                            payload.get("filename") or f"{_safe_name(org)}_карта_бизнес-процессов.docx")
    return _save_doc(doc, path)


# ===========================================================================
# ДОКУМЕНТ 2 — КАРТА РОЛЕЙ И КОНТАКТОВ · ПРЕСЕЙЛ
# ===========================================================================
ROLES_CONTACTS_SCHEMA = {
    "type": "object",
    "properties": {
        "org_name": {"type": "string", "description": "Полное наименование, напр. «АО «Рязаньавтодор»»"},
        "object_line": {"type": "string",
                        "description": "Объект: «<Компания>, ИНН ..., город (учредитель/собственник ...)»"},
        "date": {"type": "string", "description": "Дата, напр. «июнь 2026 г. (данные проверены ДД.ММ.ГГГГ)»"},
        "tldr": {"type": "array", "items": {"type": "string"},
                 "description": "TL;DR: подтверждённый ЛПР и официальный канал захода; приоритетные отделы под внедрение ИИ"},
        "org_basics": {"type": "array", "items": {"type": "string"},
                       "description": "Базовые данные: ИНН/ОГРН/КПП, адрес, учредитель, УК, численность, ОКВЭД, заказчики, финансы"},
        "leadership_table": {"type": "array", "items": {
            "type": "object",
            "properties": {"position": {"type": "string"}, "fio": {"type": "string"},
                           "responsibility": {"type": "string"}, "status": {"type": "string"}},
            "required": ["position"]},
            "description": "«1. Руководство центрального аппарата»: Должность|ФИО|Зона ответственности|Статус/источник"},
        "leadership_note": {"type": "string",
                            "description": "Примечание о неподтверждённых публично ролях (уточнять через приёмную)"},
        "departments_table": {"type": "array", "items": {
            "type": "object",
            "properties": {"block": {"type": "string"}, "contact": {"type": "string"},
                           "relevance": {"type": "string"}},
            "required": ["block", "relevance"]},
            "description": "«2. Профильные отделы — приоритет под внедрение ИИ»: Блок|Контакт|Почему релевантен и через какую боль заходить"},
        "project_institute": {"type": "string", "description": "Проектный институт / профильное подразделение (абзац, опц.)"},
        "branches_intro": {"type": "string", "description": "Вступление к таблице филиалов (опц.)"},
        "branches_table": {"type": "array", "items": {
            "type": "object",
            "properties": {"branch": {"type": "string"}, "director": {"type": "string"},
                           "phone": {"type": "string"}},
            "required": ["branch"]},
            "description": "«3. Руководители филиалов»: Филиал|Директор|Телефон (если у компании есть филиалы)"},
        "branches_note": {"type": "string", "description": "Примечание о расхождениях по числу/руководителям филиалов (опц.)"},
        "ecosystem_table": {"type": "array", "items": {
            "type": "object",
            "properties": {"entity": {"type": "string"}, "relation": {"type": "string"},
                           "person": {"type": "string"}, "contact": {"type": "string"},
                           "note": {"type": "string"}},
            "required": ["entity", "relation"]},
            "description": "«Экосистема и вертикаль принятия решений»: Организация/орган|Связь "
                           "(учредитель/курирующее ведомство/сестринская структура/комиссия/холдинг)|"
                           "Ключевое лицо|Контакт|Почему важно для захода. Критично для госкомпаний; "
                           "заполняй, если данные есть в находках"},
        "official_contacts": {"type": "array", "items": {"type": "string"},
                              "description": "«4. Официальные контакты»: приёмная/общий тел., e-mail, закупки, соцсети, график"},
        "lpr_profile_title": {"type": "string", "description": "Заголовок профиля ЛПР, напр. «Руденко С.А. — публичный деловой профиль»"},
        "lpr_profile": {"type": "string", "description": "Публичный деловой профиль ЛПР (только офиц. источники; без личных/семейных данных)"},
        "why_candidate": {"type": "array", "items": {"type": "string"},
                          "description": "Почему компания — сильный кандидат на on-premises LLM-платформу"},
        "risk_compliance": {"type": "string", "description": "Риск-факторы и комплаенс (44-ФЗ/223-ФЗ, суды, аккуратность коммуникации)"},
        "approach_plan": {"type": "array", "items": {"type": "string"},
                          "description": "«План захода»: этапы 1→4 (офиц. контакт → внутр. чемпион → тех. проработка → пилот)"},
        "threshold_signals": {"type": "array", "items": {"type": "string"},
                              "description": "Пороговые сигналы для смены тактики"},
        "disclaimers": {"type": "array", "items": {"type": "string"},
                        "description": "Оговорки по достоверности данных (что не подтверждено публично)"},
        "filename": {"type": "string", "description": "Имя файла (опц.)"},
    },
    "required": ["org_name", "tldr", "departments_table", "official_contacts"],
}


def _write_roles_contacts_docx(payload: dict, path: str = None) -> str:
    """КАРТА РОЛЕЙ И КОНТАКТОВ · ПРЕСЕЙЛ -> .docx. Возвращает путь к файлу."""
    doc, Pt, RGBColor = _new_doc()
    org = payload.get("org_name") or payload.get("company") or "Компания"

    _title_block(
        doc, Pt, RGBColor,
        "КАРТА РОЛЕЙ И КОНТАКТОВ · ПРЕСЕЙЛ",
        [org,
         "Карта ролей и деловых контактов для легитимного B2B-захода"],
        [("Подготовлено для", VENDOR),
         ("Объект", payload.get("object_line") or org),
         ("Назначение", PURPOSE_DEFAULT),
         ("Этические границы", ETHICS_DEFAULT),
         ("Дата", payload.get("date") or "")],
    )

    if payload.get("tldr"):
        _h(doc, "Краткое резюме (TL;DR)", 1)
        for b in payload["tldr"]:
            _bullet(doc, Pt, b)

    if payload.get("org_basics"):
        _h(doc, "Базовые данные организации", 1)
        for b in payload["org_basics"]:
            _bullet(doc, Pt, b)

    _h(doc, "1. Руководство центрального аппарата", 1)
    _table(doc, Pt, ["Должность", "ФИО", "Зона ответственности", "Статус / источник"],
           [[r.get("position", ""), r.get("fio", ""), r.get("responsibility", ""), r.get("status", "")]
            for r in (payload.get("leadership_table") or [])])
    if payload.get("leadership_note"):
        _para(doc, Pt, payload["leadership_note"])

    _h(doc, "2. Профильные отделы — приоритет под внедрение ИИ", 1)
    _table(doc, Pt, ["Блок", "Контактное лицо / контакт", "Почему релевантен и через какую боль заходить"],
           [[r.get("block", ""), r.get("contact", ""), r.get("relevance", "")]
            for r in (payload.get("departments_table") or [])])
    if payload.get("project_institute"):
        _h(doc, "Профильное подразделение / проектный институт", 2)
        _para(doc, Pt, payload["project_institute"])

    if payload.get("ecosystem_table"):
        _h(doc, "Экосистема и вертикаль принятия решений", 1)
        _table(doc, Pt, ["Организация / орган", "Связь", "Ключевое лицо", "Контакт", "Почему важно"],
               [[r.get("entity", ""), r.get("relation", ""), r.get("person", ""),
                 r.get("contact", ""), r.get("note", "")]
                for r in payload["ecosystem_table"]])

    if payload.get("branches_table"):
        _h(doc, "3. Руководители филиалов", 1)
        if payload.get("branches_intro"):
            _para(doc, Pt, payload["branches_intro"])
        _table(doc, Pt, ["Филиал", "Директор", "Телефон"],
               [[r.get("branch", ""), r.get("director", ""), r.get("phone", "")]
                for r in payload["branches_table"]])
        if payload.get("branches_note"):
            _para(doc, Pt, payload["branches_note"])

    _h(doc, "4. Официальные контакты компании", 1)
    for b in (payload.get("official_contacts") or []):
        _bullet(doc, Pt, b)

    if payload.get("lpr_profile") or payload.get("why_candidate") or payload.get("risk_compliance"):
        _h(doc, "Дополнительный контекст", 1)
        if payload.get("lpr_profile"):
            _h(doc, payload.get("lpr_profile_title") or "Публичный деловой профиль ЛПР", 2)
            _para(doc, Pt, payload["lpr_profile"])
        if payload.get("why_candidate"):
            _h(doc, "Почему компания — сильный кандидат на on-premises LLM-платформу", 2)
            for b in payload["why_candidate"]:
                _bullet(doc, Pt, b)
        if payload.get("risk_compliance"):
            _h(doc, "Риск-факторы и комплаенс", 2)
            _para(doc, Pt, payload["risk_compliance"])

    if payload.get("approach_plan"):
        _h(doc, "План захода (рекомендации)", 1)
        for b in payload["approach_plan"]:
            _para(doc, Pt, b)
    if payload.get("threshold_signals"):
        _h(doc, "Пороговые сигналы для смены тактики", 2)
        for b in payload["threshold_signals"]:
            _bullet(doc, Pt, b)

    if payload.get("disclaimers"):
        _h(doc, "Оговорки по достоверности данных", 1)
        for b in payload["disclaimers"]:
            _bullet(doc, Pt, b)

    if not path:
        os.makedirs(DOWNLOADS, exist_ok=True)
        path = os.path.join(DOWNLOADS,
                            payload.get("filename") or f"{_safe_name(org)}_карта_ролей_и_контактов_пресейл.docx")
    return _save_doc(doc, path)


# ===========================================================================
# ИНСТРУМЕНТЫ 2–3 — сохранение двух пресейл-документов в .docx
# ===========================================================================
@tool(
    "save_process_map_docx",
    "Сохранить КАРТУ БИЗНЕС-ПРОЦЕССОВ в .docx: профиль → as-is процессы → боли → "
    "ИИ-решение → to-be, сводная таблица, дорожная карта, специфика 44/223-ФЗ, "
    "оговорки. Ссылки/контакты — только проверенные.",
    PROCESS_MAP_SCHEMA,
)
async def save_process_map_docx(args):
    path = await anyio.to_thread.run_sync(_write_process_map_docx, args)
    return {"content": [{"type": "text", "text": f"Карта бизнес-процессов сохранена: {path}"}]}


@tool(
    "save_roles_contacts_docx",
    "Сохранить КАРТУ РОЛЕЙ И КОНТАКТОВ · ПРЕСЕЙЛ в .docx: ЛПР и руководство, профильные "
    "отделы под внедрение ИИ, экосистема/вертикаль принятия решений, филиалы, официальные "
    "контакты, план захода, оговорки. Только официальные публичные контакты; "
    "персональные/семейные данные не включать.",
    ROLES_CONTACTS_SCHEMA,
)
async def save_roles_contacts_docx(args):
    path = await anyio.to_thread.run_sync(_write_roles_contacts_docx, args)
    return {"content": [{"type": "text", "text": f"Карта ролей и контактов сохранена: {path}"}]}


research_server = create_sdk_mcp_server(
    name="research", version="3.0.0",
    tools=[deep_research, save_process_map_docx, save_roles_contacts_docx],
)


# ===========================================================================
# STAGE 1 — ТРИАЖ: грязный текст -> JSON-handle компании
# ===========================================================================
STAGE1_SYSTEM = (
    "Ты — агент идентификации компаний для российского B2B-ресёрч-инструмента.\n"
    "На входе грязный текст (описание / название / ИНН). Определи ОДНУ целевую\n"
    "компанию и аспекты для пресейла. ИНН РФ = 10 цифр (юрлицо) или 12 (ИП).\n"
    "Верни ТОЛЬКО JSON без прозы и без ```-ограждений, ровно с ключами:\n"
    '  {"company_name": str, "inn": str|null, "aspects": [str,...], '
    '"confidence": float, "notes": str}\n'
    'Если ИНН явно есть в тексте — извлеки его. Если нет — null.'
)


def _extract_json(text: str) -> dict:
    """Устойчиво вытащить первый JSON-объект из ответа модели."""
    if not text:
        raise ValueError("пустой вывод stage-1")
    fenced = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, re.DOTALL)
    candidate = fenced.group(1) if fenced else text
    try:
        return json.loads(candidate)
    except json.JSONDecodeError:
        pass
    brace = re.search(r"\{.*\}", candidate, re.DOTALL)
    if not brace:
        raise ValueError(f"JSON не найден в выводе stage-1:\n{text}")
    return json.loads(brace.group(0))


async def run_stage1_triage(raw_request: str) -> dict:
    options = ClaudeAgentOptions(
        model=FAST_MODEL,
        system_prompt=STAGE1_SYSTEM,
        max_turns=1,
        allowed_tools=[],
        setting_sources=[],
    )
    text = ""
    async for message in query(prompt=raw_request, options=options):
        if isinstance(message, AssistantMessage):
            for block in message.content:
                if isinstance(block, TextBlock):
                    text += block.text
        elif isinstance(message, ResultMessage):
            if message.result:
                text = message.result

    handle = _extract_json(text)
    handle.setdefault("company_name", "")
    handle.setdefault("inn", None)
    handle.setdefault("aspects", [])
    if not handle["company_name"] and not handle.get("inn"):
        raise ValueError(f"stage 1 не идентифицировал компанию: {handle}")
    return handle


# ===========================================================================
# STAGE 2 (ОСНОВНОЙ РЕЖИМ) — РЕСЁРЧ -> ДВА пресейл-документа .docx
# ===========================================================================
PRESALE_SYSTEM = """\
Ты — старший пресейл-аналитик и бизнес-аналитик с навыками OSINT. Готовишь ДВА
пресейл-документа по ОДНОЙ компании под продажу корпоративной on-premises
LLM-платформы. Платформа: RAG-база знаний (семантический поиск по нормативке/
документации со ссылкой на источник), автономные ИИ-агенты (многошаговые задачи:
мониторинг, проверка, генерация документов), ИИ-Коуч «Наставник» (онбординг и
тиражирование практик). Ключевое УТП — on-premises: данные не уходят в облако
(152-ФЗ, импортозамещение, ИБ).

ПОРЯДОК РАБОТЫ (ресёрч делаешь ОДИН раз, оба документа — из одних находок):
1) deep_research(company_name, inn) — ЗОВИ ПЕРВЫМ и ОДИН раз. Это уже НЕ только
   официальная база: движок сам открывает САЙТ компании (филиалы и их директора+
   телефоны, руководство, контакты, СОЦСЕТИ, проектный институт), ЕИС/госзакупки по
   ИНН, суды/СМИ и hh.ru, гоняет петлю ЦЕЛЕВОГО добора под пустые ячейки таблиц и
   возвращает источникованный документ: markdown + структурный JSON (у КАЖДОЙ строки
   source URL), список РЕАЛЬНО ОТКРЫТЫХ страниц и блок «Полнота целевых таблиц».
   ОПИРАЙСЯ на эти находки как на основу обоих документов — особенно таблицу филиалов,
   контакты закупок и соцсети (раньше их теряли).
2) WebSearch + WebFetch — доверка и закрытие ЛЮБЫХ существенных пробелов, НЕ ограничивайся
   фиксированным списком: остаточные пробелы из блока «Полнота целевых таблиц» (контакт
   ИТ/цифровизации, ФИО замов, контактное лицо закупок), ВЕРТИКАЛЬ принятия решений
   (учредитель/собственник, курирующее ведомство/министерство, сестринские структуры,
   комиссии/советы по теме ИИ), раздел «Команда/Руководство» на сайте компании, свежие
   назначения и профильные программы в деловых СМИ. Не дублируй то, что deep_research уже
   принёс. Проверяй спорные ссылки WebFetch'ем; битые/выдуманные НЕ включай. Контакты бери
   ТОЛЬКО официальные (сайт, извещения о закупках, ЕГРЮЛ). Персональные/семейные данные НЕ собирай.
3) Вызови save_process_map_docx — КАРТА БИЗНЕС-ПРОЦЕССОВ:
   org_name; tldr; key_findings; profile_table (Параметр|Значение|Источник);
   financials; structure_qc; contracts_table; processes — 6–9 процессов компании,
   по каждому as_is → pains → ai_solution (RAG/ИИ-агент/ИИ-Коуч и что делает) →
   to_be (измеримо); summary_table (Процесс|Боль|Продукт/решение|Эффект);
   roadmap (этап 0 демо → 1 пилот «быстрые победы» → 2 масштаб → 3 зрелость);
   gov_sale (путь продажи через конкурентную процедуру, ЛПР и функц. заказчики,
   on-prem УТП, обоснование НМЦК, триггерные события); recommendations; disclaimers.
4) Вызови save_roles_contacts_docx — КАРТА РОЛЕЙ И КОНТАКТОВ · ПРЕСЕЙЛ:
   org_name; object_line; tldr; org_basics; leadership_table (Должность|ФИО|Зона
   ответственности|Статус/источник — неподтверждённое помечай); leadership_note;
   departments_table (Блок|Контакт|Почему релевантен и через какую боль заходить —
   расставь приоритет под внедрение ИИ, явно отметь недостающий контакт ИТ/
   цифровизации); ecosystem_table (Организация|Связь|Ключевое лицо|Контакт|Почему
   важно — вертикаль принятия решений: учредитель, курирующее ведомство, сестринские
   структуры, комиссии; переноси из блока «Экосистема» находок, для госкомпаний
   обязательно, если данные есть); project_institute; branches_table (если есть филиалы);
   official_contacts (приёмная, e-mail, закупки, соцсети, график); lpr_profile
   (только офиц. источники); why_candidate; risk_compliance; approach_plan
   (этапы 1→4); threshold_signals; disclaimers.

ПРИНЦИПЫ:
- Только открытые данные. Каждый существенный факт — с источником; разделяй
  ФАКТ / ОЦЕНКУ / ГИПОТЕЗУ. Эффекты ИИ помечай как отраслевые оценки, не измерения.
- Бухгалтерскую выручку (стр. 2110) и обороты/объём контрактов НЕ путать —
  указывай метрику, год, источник. Расхождения по первому лицу/числу филиалов
  фиксируй в оговорках.
- Заход — строго легитимный: официальные каналы и конкурентная процедура, без
  обхода закупок через личные связи. Никаких выдуманных цифр, ФИО и ссылок.
- ДИСЦИПЛИНА ДОЗАБОРА (обязательно):
  (A1) Ячейку (директор филиала, контакт отдела, ФИО руководителя и т.п.) можно
       оставить «не подтверждено» ТОЛЬКО если соответствующий источник РЕАЛЬНО открыт
       (он есть в списке «Реально открытые источники» из deep_research или ты сам
       открыл его WebFetch'ем) и факта там нет. ЛЮБОЕ отрицание — со ссылкой на
       открытую страницу/выдачу ЕИС. Если deep_research уже принёс ФИО+телефон
       филиала — переноси их, НЕ пиши «не подтверждено». Запрещено и выдумывать, и
       уверенно отрицать без открытой страницы.
  (A2) ПЕРЕД сохранением — самопроверка полноты: пройди по структурному JSON находок
       и блоку «Полнота целевых таблиц». Все принесённые филиалы (директор+телефон),
       соцсети и контакты закупок ДОЛЖНЫ попасть в документы. Если в JSON строка есть,
       а в таблице её нет — добавь её. Пустые ячейки в leadership/branches/departments
       оставляй только по правилу A1, с источником-ссылкой.
Оба документа ОБЯЗАТЕЛЬНО сохрани (вызови ОБА инструмента). В конце — краткое
резюме: что сохранено и топ-3 точки внедрения."""


# ===========================================================================
# СИСТЕМНЫЙ ПРОМПТ ТРЕТЬЕГО ДЕЛИВЕРАБЛА — 3-слайдовая презентация .pptx
#   Здесь форматы документов проекта живут в CRA, поэтому промпт презентации —
#   тоже тут, рядом с PRESALE_SYSTEM. В отличие от двух .docx (детерминированные
#   рендереры), презентацию делает АГЕНТНАЯ сессия через официальный скилл pptx
#   (code-execution/Bash), а НЕ python-рендерер. Текст промпта — VERBATIM.
# ===========================================================================
PRESENTATION_SYSTEM = """\
КОНТЕКСТ И РОЛЬ
Ты — дизайнер презентаций и бизнес-аналитик. Я отвечаю за коммерцию и продукт платформы Telepath в АО «ЦИТ РТ» (Центр информационных технологий Республики Татарстан). Telepath — корпоративная LLM-платформа (ИИ в защищённом on-premises контуре): RAG-базы знаний, автономные ИИ-агенты, ИИ-Коуч. Я готовлю клиентскую презентацию для потенциального заказчика.

ВХОДНЫЕ ДАННЫЕ (заполняю под каждого клиента)
- НАЗВАНИЕ КОМПАНИИ-ЗАКАЗЧИКА:
- ЧЕМ ЗАНИМАЕТСЯ / ОТРАСЛЬ:
- ГЛАВНАЯ БОЛЬ (если знаю):

ЗАДАЧА
Сделай готовую к отправке презентацию .pptx (16:9, widescreen), РОВНО 3 слайда, в фирменном стиле сайта https://citrt.ru/. Её читает сам заказчик и в финале должен захотеть записаться на демо. Выполни за один проход: сгенерируй файл, отрендери каждый слайд в картинку, проверь отсутствие переполнений и налезаний, исправь и только потом отдай .pptx.
Слайды 1 и 2 неизменны (о нас и о нашем эксперте). Слайд 3 — оффер — адаптируй под компанию-заказчика из «ВХОДНЫХ ДАННЫХ».

ВЛОЖЕНИЯ (в этом же запросе)
1) Логотип ЦИТ РТ (PNG). 2) Фото Булата Замалиева (PNG) — кадрируй в портрет, убери посторонний фон/цветовые пятна по краям, скругли углы.

ФИРМЕННЫЙ СТИЛЬ (палитра из логотипа/сайта)
- Голубой #0090C0 (основной), лаймовый #8CC63F (акцент, как «С» в логотипе), тёмно-бирюзовый #0B2B3C (тёмные секции), белый (контент). Доп.: тёмный голубой #0A6E96, серо-стальной #55687A.
- Чистый современный дизайн: карточки со скруглением и мягкой тенью; иконки в кружках как единый мотив; БЕЗ полосок-акцентов и подчёркиваний под заголовками; шрифт без засечек (Arial).
- Логотип: на тёмных слайдах — на белой скруглённой подложке с тонкой голубой обводкой; на белом слайде — в углу без подложки.

СЛАЙД 1 — «О компании АО ЦИТ РТ» (тёмно-бирюзовый фон) — НЕИЗМЕНЕН
- Надзаголовок «РАЗРАБОТЧИК ПЛАТФОРМЫ TELEPATH»; заголовок «АО «ЦИТ РТ»»; подзаголовок: государственный Центр информационных технологий Республики Татарстан, разработчик Telepath.
- Слева — 4 карточки-преимущества (кружок-иконка + заголовок + короткое описание) одной колонкой: Суверенность данных (on-premises, 152-ФЗ); Импортозамещение (российская разработка, автономный контур); Модель «гос-к-гос» (госкомпания РТ — доверенный партнёр для предприятий с госучастием); Отраслевой опыт (энергетика, нефтегаз, ритейл, госсектор).
- Справа — крупный логотип на белой панели.
- В подвале мелким: ИНН 1655505808, ОГРН 1241600056829, обслуживание в ПАО «АК БАРС» Банк.

СЛАЙД 2 — «Булат Замалиев», руководитель направления ИИ (тёмно-бирюзовый фон) — НЕИЗМЕНЕН
- Слева — его фото в голубой рамке. Справа: надзаголовок «КОМАНДА TELEPATH»; имя «Булат Замалиев»; должность «Руководитель направления ИИ, АО «ЦИТ РТ»».
- Блок «КЛЮЧЕВЫЕ ДОСТИЖЕНИЯ» — только проверенные факты (ничего не выдумывай; при сомнении проверь по авторитетным источникам: Forbes, «Российская газета», «Коммерсантъ», сайты вузов). Используй именно эти формулировки:
  • Уполномоченный по технологиям ИИ в Республике Татарстан — первый в России (с 2019).
  • Номинант рейтинга Forbes «30 до 30» (2020), категория «Управление». (Именно НОМИНАНТ/лонг-лист, НЕ «победитель».)
  • Победитель всероссийского хакатона «Цифровой прорыв» (2019), грант 3 млн ₽.
  • Идеолог ИИ-сервиса «Госпромпт» для госслужащих и проекта «Цифровая деревня».
- Внизу зелёная плашка-связка: «→ Проведёт демонстрацию Telepath для вашей команды».
- ДОСТОВЕРНОСТЬ: публичные источники подтверждают его как «Уполномоченного по технологиям ИИ при Минцифры РТ»; связь с ЦИТ РТ как руководителя направления публично не подтверждена. Поставь должность как указано, но отдельным замечанием предупреди меня об этом расхождении и предложи проверяемую альтернативу.

СЛАЙД 3 — Оффер «Серьёзная проблема — простое решение» (белый фон) — АДАПТИРУЙ ПОД ЗАКАЗЧИКА
- Надзаголовок «ПЕРСОНАЛЬНО ДЛЯ [НАЗВАНИЕ КОМПАНИИ]»; заголовок «Серьёзная проблема — простое решение»; логотип в углу.
- Сформулируй главную боль заказчика исходя из исследований прошлых агентов.
- Две колонки:
  • Слева (нейтральная серая карточка, серый кружок-иконка) — ГЛАВНАЯ ПРОБЛЕМА: [ёмкий заголовок боли]. 3–4 коротких пункта, конкретных под отрасль заказчика (где уходит время специалистов, что делается вручную, почему ошибки дороги, где разрозненность).
  • Справа (светло-голубая карточка, зелёный кружок-иконка робота) — РЕШЕНИЕ: платформа Telepath. ИИ-агенты берут рутину на себя; RAG-база даёт мгновенный ответ по нормативке/регламентам; ускорение профильной операции с часов до минут; on-premises — данные не покидают предприятие.
- Полоса метрик (4 плитки, 3 голубые + 1 зелёная), подпиши под отрасль заказчика: «≈8×» (быстрее профильная операция); «Минуты» (вместо часов на типовую задачу); «Секунды» (ответ по нормативке/регламентам через базу знаний); «On-prem» (данные в вашем контуре). НЕ используй «24/7».
- Внизу тёмно-бирюзовая плашка-призыв: «Назначьте бесплатное демо Telepath» + «Покажем на примере ваших реальных задач — без обязательств и доступа к вашим данным» + контакты зелёным: [телефон], [email], сайт citrt.ru.

ТРЕБОВАНИЯ К КАЧЕСТВУ
- Ничего не вылезает за границы плашек/слайдов — проверь рендер всех 3 слайдов и поправь до выдачи.
- Контакты, которых у тебя нет, оставляй явными полями-заполнителями (например, [телефон]).
- Факты о людях/компании — только подтверждённые; сомнительное помечай отдельным замечанием.
- Отдай один готовый файл .pptx."""


async def run_dossier(handle: dict) -> str:
    options = ClaudeAgentOptions(
        model=MODEL,
        system_prompt=PRESALE_SYSTEM,
        mcp_servers={"research": research_server},
        allowed_tools=[
            "mcp__research__deep_research",
            "mcp__research__save_process_map_docx",
            "mcp__research__save_roles_contacts_docx",
            "WebSearch", "WebFetch",
        ],
        # запрещаем побочные эффекты ФС; данные пишут только save_*_docx
        disallowed_tools=["Bash", "Edit", "Write", "NotebookEdit"],
        permission_mode="bypassPermissions",   # headless: без зависаний на аппруве инструментов
        setting_sources=[],
        max_turns=80,                          # глубокий ресёрч + проверка ссылок + ДВА документа
    )

    aspects_csv = ", ".join(handle.get("aspects") or
                            ["профиль, процессы as-is, точки внедрения ИИ, оргструктура, ЛПР, контакты, СМИ за 5 лет"])
    handoff = (
        "Подготовь ДВА пресейл-документа (карта бизнес-процессов + карта ролей "
        "и контактов пресейл) и сохрани ОБА в .docx. Значения для инструментов:\n"
        f"  company_name = {handle['company_name']!r}\n"
        f"  inn          = {handle['inn'] or ''!r}\n"
        f"  aspects      = {aspects_csv!r}\n"
        "Следуй порядку работы из системного промпта. Ссылки/контакты — только проверенные."
    )

    summary = ""
    async with ClaudeSDKClient(options=options) as client:
        await client.query(handoff)
        async for message in client.receive_response():
            if isinstance(message, AssistantMessage):
                for block in message.content:
                    if isinstance(block, TextBlock):
                        summary += block.text
                    elif isinstance(block, ToolUseBlock):
                        nm = getattr(block, "name", "")
                        print(f"  → {nm}")
            elif isinstance(message, ResultMessage):
                if message.result:
                    summary = message.result
                print(f"  [пресейл] cost=${message.total_cost_usd} {message.subtype}")
    return summary


# Старый режим (краткий отчёт по контактам/официалке, без СМИ и без .docx) — по флагу --contacts
async def run_contacts(handle: dict) -> str:
    payload = await anyio.to_thread.run_sync(
        research_company, handle["company_name"], handle.get("inn") or "",
        ", ".join(handle.get("aspects") or []))
    return _render_official(payload)


async def run_pipeline(raw_request: str, mode: str = "dossier") -> str:
    print("=== STAGE 1: триаж ===")
    handle = await run_stage1_triage(raw_request)
    print("  handle:", json.dumps(handle, ensure_ascii=False))
    if mode == "contacts":
        print("=== STAGE 2: контакты/официалка ===")
        return await run_contacts(handle)
    print("=== STAGE 2: ресёрч -> карта бизнес-процессов + карта ролей и контактов (.docx) ===")
    return await run_dossier(handle)


async def main():
    argv = sys.argv[1:]
    mode = "dossier"
    if argv and argv[0] == "--contacts":
        mode, argv = "contacts", argv[1:]
    cli = " ".join(argv).strip()
    try:
        if cli:                                   # одноразовый режим
            print("\n" + await run_pipeline(cli, mode))
            return
        print("Пресейл-агент по компаниям. Пустая строка или 'exit' — выход.")
        while True:                               # интерактивный режим
            try:
                req = input("\n> ").strip()
            except (EOFError, KeyboardInterrupt):
                break
            if not req or req.lower() in ("exit", "quit", "выход"):
                break
            print("\n" + await run_pipeline(req, mode))
    except (ValueError, json.JSONDecodeError) as e:
        print(f"Ошибка разбора stage-1: {e}")


if __name__ == "__main__":
    anyio.run(main)
