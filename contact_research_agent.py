# -*- coding: utf-8 -*-
r"""
Второй агент-ресёрчер (КОНТАКТЫ И ТОЧКИ ВХОДА) на Claude Agent SDK.

Опирается на JSON лида + карту бизнес-процессов первого агента
(company_research_agent) и через deepsearch (официальная база + широкий
WebSearch/WebFetch) собирает:

  - key_contacts : топ точки входа (кто; чем релевантен; как выйти/через кого)
  - roles        : должность, зона ответственности, контакты, ссылки на профили
  - branches     : филиалы и их руководители
  - sources      : ТОЛЬКО проверенные WebFetch'ем ссылки

Инструменты агента:
  - mcp__contacts__deep_research      : официальная база по ИНН (переиспользуем
        company_research_agent.deep_research — Dadata: ЛПР/филиалы, Checko: контакты)
  - WebSearch / WebFetch (встроенные) : раздел сайта «команда/руководство», hh.ru,
        СМИ/интервью, подписанты госзакупок, профессиональные профили
        КАЖДАЯ ссылка (контакты и sources) проверяется WebFetch'ем (анти-галлюцинация)
  - mcp__contacts__save_contacts_docx : рендер отчёта по контактам в .docx

Запуск (одиночная отладка по одной компании):
  py contact_research_agent.py "АО Рязаньавтодор ИНН 6234065445"
  py contact_research_agent.py            # без аргумента — интерактивный режим

Зависимости:  pip install claude-agent-sdk python-docx
"""
import json
import os
import re
import sys

import anyio
from claude_agent_sdk import (
    create_sdk_mcp_server,  # in-process MCP-сервер (без subprocess)
    tool,                   # @tool — in-process MCP-инструмент
    ClaudeAgentOptions,     # опции (поля snake_case)
    ClaudeSDKClient,        # многоходовый клиент (context manager)
    AssistantMessage, TextBlock, ToolUseBlock, ResultMessage,
)

# консоль Windows = cp1251 и роняет вывод на ₽/кириллице -> принудительно utf-8
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# Все модули скилла лежат в этом же каталоге — добавляем его в sys.path.
SCRIPTS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPTS)

# Первый агент: переиспользуем его инструмент официальной базы и хелперы рендера.
import company_research_agent as CRA  # noqa: E402

MODEL = "opus"          # ресёрч контактов (синтез + проверка ссылок)
FAST_MODEL = "sonnet"   # триаж грязного ввода

DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")


# ===========================================================================
# ИНСТРУМЕНТ — рендер отчёта по контактам в .docx (многостраничный) в «Загрузки»
# ===========================================================================
CONTACTS_SCHEMA = {
    "type": "object",
    "properties": {
        "company": {"type": "string", "description": "Заголовок: «Контакты и точки входа — Компания»"},
        "subtitle": {"type": "string", "description": "ИНН · ОГРН · отрасль · город"},
        "summary": {"type": "string",
                    "description": "Как лучше всего зайти (2–4 строки): оптимальная точка входа и канал"},
        "key_contacts": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "name": {"type": "string", "description": "ФИО (или роль, если ФИО не найдено)"},
                "position": {"type": "string", "description": "Должность"},
                "why_relevant": {"type": "string", "description": "Чем релевантен / почему это точка входа под наш оффер"},
                "how_to_reach": {"type": "string", "description": "Как выйти: канал (email/звонок/площадка) и через кого"},
                "priority": {"type": "string", "description": "Приоритет: высокий / средний / низкий"}},
            "required": ["name", "why_relevant", "how_to_reach"]},
            "description": "Топ точки входа: кто; чем релевантен; как выйти"},
        "roles": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "fio": {"type": "string", "description": "ФИО"},
                "position": {"type": "string", "description": "Должность"},
                "responsibility": {"type": "string", "description": "Зона ответственности"},
                "contacts": {"type": "string", "description": "Контакты (тел/email), если есть"},
                "profiles": {"type": "string", "description": "Ссылки на профили (через пробел/запятую), только проверенные"}},
            "required": ["position", "responsibility"]},
            "description": "Роли: должность, зона ответственности, контакты, ссылки на профили"},
        "branches": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "name": {"type": "string", "description": "Название/обозначение филиала"},
                "location": {"type": "string", "description": "Город/регион"},
                "head": {"type": "string", "description": "ФИО руководителя филиала"},
                "head_position": {"type": "string", "description": "Должность руководителя"},
                "contacts": {"type": "string", "description": "Контакты филиала (если есть)"}},
            "required": ["name"]},
            "description": "Филиалы и их руководители"},
        "sources": {"type": "array", "items": {
            "type": "object",
            "properties": {"title": {"type": "string"}, "url": {"type": "string"}},
            "required": ["title", "url"]},
            "description": "ТОЛЬКО проверенные WebFetch'ем ссылки-источники"},
        "filename": {"type": "string", "description": "Имя файла (опц.)"},
    },
    "required": ["company"],
}


def _write_contacts_docx(payload: dict) -> str:
    """Собрать .docx-отчёт по контактам (key_contacts/roles/branches/sources). Возвращает путь."""
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor

    GRAY = RGBColor(0x5A, 0x5A, 0x5A)
    LINK = RGBColor(0x15, 0x4F, 0x9C)
    doc = Document()

    # компактные поля и базовый шрифт
    sec = doc.sections[0]
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(1.2))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    pf = normal.paragraph_format
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(11)
        return p

    def bullet(text):
        # жирная лид-метка до первого ':' либо '—', остальное обычным
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        m = re.match(r"^(.*?[:—])(\s*)(.*)$", text, re.DOTALL)
        if m:
            p.add_run(m.group(1)).bold = True
            if m.group(3):
                p.add_run(" " + m.group(3))
        else:
            p.add_run(text)
        return p

    def person_card(title, meta, rows):
        """Карточка контакта/роли: жирный заголовок + серая мета + помеченные строки."""
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(1)
        r = cp.add_run(title or "—")
        r.bold = True
        r.font.size = Pt(10)
        if meta:
            mr = cp.add_run("  · " + meta)
            mr.font.size = Pt(8.5)
            mr.font.color.rgb = GRAY
        for lbl, val in rows:
            if val:
                bullet(f"{lbl}: {val}")

    # шапка
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(1)
    tr = title.add_run(payload.get("company") or "Контакты и точки входа")
    tr.bold = True
    tr.font.size = Pt(13)
    if payload.get("subtitle"):
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)
        sr = sp.add_run(payload["subtitle"])
        sr.font.size = Pt(8.5)
        sr.font.color.rgb = GRAY

    if payload.get("summary"):
        sm = doc.add_paragraph()
        sm.paragraph_format.space_before = Pt(2)
        sm.paragraph_format.space_after = Pt(4)
        smr = sm.add_run(payload["summary"])
        smr.italic = True
        smr.font.size = Pt(9.5)

    key_contacts = payload.get("key_contacts") or []
    if key_contacts:
        heading("1. Ключевые точки входа")
        for it in key_contacts:
            meta = " · ".join(x for x in (it.get("position"), it.get("priority")) if x)
            person_card(it.get("name"), meta, [
                ("Чем релевантен", it.get("why_relevant")),
                ("Как выйти", it.get("how_to_reach")),
            ])

    roles = payload.get("roles") or []
    if roles:
        heading("2. Роли и зоны ответственности")
        for it in roles:
            person_card(it.get("fio") or it.get("position") or "—",
                        it.get("position") if it.get("fio") else "",
                        [
                            ("Зона ответственности", it.get("responsibility")),
                            ("Контакты", it.get("contacts")),
                            ("Профили", it.get("profiles")),
                        ])

    branches = payload.get("branches") or []
    if branches:
        heading("3. Филиалы и их руководители")
        for it in branches:
            meta = it.get("location") or ""
            head = " — ".join(x for x in (it.get("head"), it.get("head_position")) if x)
            person_card(it.get("name"), meta, [
                ("Руководитель", head),
                ("Контакты", it.get("contacts")),
            ])

    sources = payload.get("sources") or []
    if sources:
        heading("4. Источники (проверенные ссылки)")
        for i, s in enumerate(sources, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.add_run(f"{i}. {s.get('title', '')} — ")
            link = p.add_run(s.get("url", ""))
            link.font.color.rgb = LINK
            link.font.size = Pt(8.5)

    os.makedirs(DOWNLOADS, exist_ok=True)
    base = payload.get("company", "") or "Компания"
    fname = payload.get("filename") or f"Контакты_{CRA._safe_name(base)}.docx"
    if not fname.lower().endswith(".docx"):
        fname += ".docx"
    path = os.path.join(DOWNLOADS, fname)
    doc.save(path)
    return path


@tool(
    "save_contacts_docx",
    "Сохранить готовый отчёт по контактам и точкам входа в .docx (многостраничный) в папку "
    "«Загрузки». Блоки: key_contacts (точки входа), roles (роли/зоны ответственности), "
    "branches (филиалы и руководители), sources. Ссылки в profiles/sources — только "
    "проверенные WebFetch'ем.",
    CONTACTS_SCHEMA,
)
async def save_contacts_docx(args):
    path = await anyio.to_thread.run_sync(_write_contacts_docx, args)
    return {"content": [{"type": "text", "text": f"Контакты сохранены: {path}"}]}


contacts_server = create_sdk_mcp_server(
    name="contacts", version="1.0.0",
    tools=[CRA.deep_research, save_contacts_docx],
)


# ===========================================================================
# СИСТЕМНЫЙ ПРОМПТ — поиск контактов и точек входа
# ===========================================================================
CONTACTS_SYSTEM = """\
Ты — OSINT-аналитик B2B-продаж. Твоя задача — по компании собрать карту ЛИЦ и
ТОЧЕК ВХОДА для outbound-продажи вендора ИИ-решений (on-prem LLM, ИИ-агенты, RAG).
На входе: реквизиты компании и КАРТА БИЗНЕС-ПРОЦЕССОВ (AS-IS → точки внедрения ИИ →
TO-BE) от первого агента — используй её, чтобы понять, КАКИЕ функции и роли наиболее
релевантны под наш оффер (например: ИТ/цифровизация, операционный директор,
руководитель направления, где сосредоточена «боль» из карты процессов).

ПОРЯДОК РАБОТЫ (deepsearch — официальная база + широкий веб):
1) Вызови deep_research(company_name, inn, aspects) — официальная база: карточка
   ЕГРЮЛ/ЕГРИП (ЛПР/ОКВЭД/адрес/филиалы — Dadata), контакты (Checko). Это каркас:
   первое лицо и филиалы.
2) Через WebSearch + WebFetch добери из ОТКРЫТЫХ источников людей и точки входа:
   - сайт компании, разделы «Команда / Руководство / Контакты / Филиалы»;
   - hh.ru (вакансии → ответственные за функции, оргструктура, стек);
   - СМИ/интервью/пресс-релизы (имена топ-менеджеров, спикеры);
   - госзакупки/тендеры (подписанты, контактные лица закупок);
   - профессиональные профили и отраслевые каталоги.
   НЕ ограничивайся официальными базами — они дают только ЛПР и филиалы.
3) КАЖДУЮ ссылку, которую кладёшь в profiles или sources, ОБЯЗАТЕЛЬНО открой через
   WebFetch и убедись, что страница реальна (не 404) и именно про этого человека/
   компанию. Битые, непроверенные и выдуманные ссылки НЕ включай.
4) Сформируй и вызови save_contacts_docx со ВСЕМИ блоками:
   company      — «Контакты и точки входа — Компания»;
   subtitle     — ИНН · ОГРН · отрасль · город;
   summary      — как лучше всего зайти (оптимальная точка входа и канал, 2–4 строки);
   key_contacts — ТОП точки входа: name (ФИО или роль), position, why_relevant
                  (чем релевантен — привяжи к процессам/болям из карты), how_to_reach
                  (как выйти: канал email/звонок/площадка и через кого), priority;
   roles        — роли: fio, position, responsibility (зона ответственности),
                  contacts (тел/email если есть), profiles (проверенные ссылки);
   branches     — филиалы: name, location, head (ФИО руководителя), head_position,
                  contacts;
   sources      — ТОЛЬКО проверенные WebFetch'ем ссылки (title + url).

ПРИНЦИПЫ:
- Опирайся ТОЛЬКО на открытые данные и официальные базы. Каждый человек/факт — с
  источником. Разделяй ФАКТ / ОЦЕНКУ / ГИПОТЕЗУ.
- Никаких выдуманных ФИО, должностей, контактов и ссылок. Нет данных — пиши «не
  найдено в открытых источниках» и не выдумывай.
- Приоритизируй точки входа по релевантности карте бизнес-процессов и реализуемости
  контакта (есть ли прямой канал).
После сохранения дай краткое резюме: куда сохранён файл и топ-1 рекомендуемая точка входа."""


def _lead_brief(lead: dict) -> str:
    """Краткие реквизиты лида для хэндоффа агенту."""
    def g(k):
        return str(lead.get(k) or "").strip()

    lpr = (g("contact_person") + " " + g("_lpr_post")).strip()
    parts = [
        f"company_name = {g('name')!r}",
        f"inn          = {g('_inn')!r}",
        f"ogrn         = {g('_ogrn')!r}",
        f"отрасль      = {(g('niche') or g('_industry'))!r}",
        f"регион       = {g('_region')!r}",
        f"адрес        = {g('_address')!r}",
        f"сайт         = {g('website')!r}",
        f"известный ЛПР = {lpr!r}",
    ]
    return "\n".join("  " + p for p in parts)


def _build_handoff(lead: dict, biz_context: str = "") -> str:
    """Текст хэндоффа агенту-2: реквизиты лида + карта бизнес-процессов агента-1."""
    ctx = (biz_context or "").strip()
    ctx_block = (
        "\nКАРТА БИЗНЕС-ПРОЦЕССОВ (от первого агента) — используй для приоритизации точек входа:\n"
        + ctx + "\n"
    ) if ctx else "\n(Карта бизнес-процессов не передана — опирайся на реквизиты и веб-ресёрч.)\n"
    return (
        "Собери контакты и точки входа по компании и сохрани отчёт в .docx. Реквизиты:\n"
        + _lead_brief(lead) + "\n"
        + ctx_block
        + "Следуй порядку из системного промпта. profiles/sources — только проверенные WebFetch'ем ссылки."
    )


async def run_contacts_agent(lead: dict, biz_context: str = "") -> str:
    """Запустить агента контактов по лиду + карте бизнес-процессов. Возвращает резюме."""
    options = ClaudeAgentOptions(
        model=MODEL,
        system_prompt=CONTACTS_SYSTEM,
        mcp_servers={"contacts": contacts_server},
        allowed_tools=[
            "mcp__contacts__deep_research",
            "mcp__contacts__save_contacts_docx",
            "WebSearch", "WebFetch",
        ],
        # запрещаем побочные эффекты ФС; данные пишет только save_contacts_docx
        disallowed_tools=["Bash", "Edit", "Write", "NotebookEdit"],
        permission_mode="bypassPermissions",   # headless: без зависаний на аппруве инструментов
        setting_sources=[],
        max_turns=60,                          # глубокий ресёрч + проверка ссылок WebFetch'ем
    )

    handoff = _build_handoff(lead, biz_context)

    summary = ""
    async with ClaudeSDKClient(options=options) as client:
        await client.query(handoff)
        async for message in client.receive_response():
            if isinstance(message, AssistantMessage):
                for block in message.content:
                    if isinstance(block, TextBlock):
                        summary += block.text
                    elif isinstance(block, ToolUseBlock):
                        print(f"  → {getattr(block, 'name', '')}")
            elif isinstance(message, ResultMessage):
                if message.result:
                    summary = message.result
                print(f"  [contacts] cost=${message.total_cost_usd} {message.subtype}")
    return summary


# ===========================================================================
# Автономный режим: триаж грязного ввода -> агент контактов (для отладки)
# ===========================================================================
async def run_pipeline(raw_request: str) -> str:
    print("=== STAGE 1: триаж ===")
    handle = await CRA.run_stage1_triage(raw_request)
    print("  handle:", json.dumps(handle, ensure_ascii=False))
    lead = {
        "name": handle.get("company_name") or "",
        "_inn": handle.get("inn") or "",
    }
    print("=== STAGE 2: контакты и точки входа (ресёрч + .docx) ===")
    return await run_contacts_agent(lead, biz_context="")


async def main():
    cli = " ".join(sys.argv[1:]).strip()
    try:
        if cli:                                   # одноразовый режим
            print("\n" + await run_pipeline(cli))
            return
        print("Агент контактов и точек входа. Пустая строка или 'exit' — выход.")
        while True:                               # интерактивный режим
            try:
                req = input("\n> ").strip()
            except (EOFError, KeyboardInterrupt):
                break
            if not req or req.lower() in ("exit", "quit", "выход"):
                break
            print("\n" + await run_pipeline(req))
    except (ValueError, json.JSONDecodeError) as e:
        print(f"Ошибка разбора stage-1: {e}")


if __name__ == "__main__":
    anyio.run(main)
