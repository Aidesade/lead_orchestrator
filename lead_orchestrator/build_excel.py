# -*- coding: utf-8 -*-
"""
Генерация Excel «База лидов» в формате шаблона D:\\лиды\\Пример_лиды.xlsx.

Колонки (лист «База лидов»):
 A №  B Компания  C Сфера/ниша  D Сайт  E Телефон  F Email
 G Контактное лицо / ЛПР  H Источник (2ГИС/СПАРК/...)
 I Боль / потребность в ИИ  J Предлагаемый оффер  K Статус  L След. шаг / дата
"""
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HEADERS = [
    "№", "Компания", "Сфера/ниша", "Сайт", "Телефон", "Email",
    "Контактное лицо / ЛПР", "Источник (2ГИС/СПАРК/...)",
    "Боль / потребность в ИИ", "Предлагаемый оффер", "Статус", "След. шаг / дата",
    # доказательная база ЛПР (заполняется этапом обогащения dadata_enrich)
    "Должность ЛПР", "ИНН", "Регион", "Выручка, ₽", "Источник выручки",
    "Достоверность ЛПР",
]
LAST_COL = "R"  # последняя колонка (12 базовых + 6 обогащения)
# ширины колонок A..R (как в шаблоне; D и G добавлены для читаемости)
COL_WIDTHS = {
    "A": 5, "B": 26, "C": 24, "D": 24, "E": 16, "F": 24,
    "G": 22, "H": 20, "I": 30, "J": 28, "K": 16, "L": 20,
    "M": 20, "N": 14, "O": 22, "P": 18, "Q": 26, "R": 16,
}

_STATUS_RU = {
    "ACTIVE": "Действует", "LIQUIDATING": "Ликвидируется",
    "LIQUIDATED": "Ликвидирована", "BANKRUPT": "Банкротство",
    "REORGANIZING": "Реорганизация",
}


def _status_ru(s):
    return _STATUS_RU.get(s or "", s or "")


def _confidence_label(ld):
    t = ld.get("_match_tier")
    if not t:
        return ""
    c = ld.get("_match_confidence")
    return f"{t} ({c})" if c is not None else t


# цвет шрифта Email по типу почты: целевая зелёная/синяя, общая — оранжевая,
# служебная — серая: это техподдержка или робот, работать по ней нельзя
_EMAIL_COLOR = {
    "личная ЛПР": "FF1E7A34", "личная": "FF1E7A34", "руководитель": "FF1E7A34",
    "почта компании": "FF1F6FB2", "сотрудничество": "FF1F6FB2",
    "общая": "FFC55A11", "служебная": "FF808080",
}


def _mark_email(cell, ld):
    kind = ld.get("_email_kind", "")
    color = _EMAIL_COLOR.get(kind)
    if color:
        cell.font = Font(name="Arial", size=10, color=color,
                         bold=kind in ("личная ЛПР", "сотрудничество"),
                         italic=kind == "служебная")
    src = ld.get("_email_src", "")
    note = f"Тип почты: {kind}"
    if src:
        note += f"\nИсточник: {src}"
    if kind == "служебная":
        note += ("\n⚠ служебный ящик (техподдержка/робот) — НЕ контактный. "
                 "Взят потому, что другого адреса у компании не нашлось")
    elif not ld.get("_email_is_target"):
        note += "\n⚠ общая почта (не ЛПР) — целевой адрес на сайте не найден"
    # Проверка домена (source_girbo.check_mail): адрес, который не доставится,
    # должен быть виден до рассылки, а не после отбойника
    check = ld.get("_email_check")
    if check:
        note += f"\nПроверка домена: {check} — {ld.get('_email_note', '')}"
        if check in ("не доставится", "битый", "одноразовый"):
            cell.font = Font(name="Arial", size=10, color="FFC00000", strike=True)
    if ld.get("_email_verified"):
        note += "\n✓ подтверждена на офиц. сайте"
    cell.comment = Comment(note, "lead-finder")

TITLE_FILL = PatternFill("solid", fgColor="FF1F3864")
HEAD_FILL = PatternFill("solid", fgColor="FF2E5496")
WHITE = "FFFFFFFF"
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ZEBRA = PatternFill("solid", fgColor="FFF2F5FA")


def build(leads, industry, out_path):
    """leads: список dict с ключами name, niche, website, phone, email,
    contact_person, source, pain, offer, status, next_step."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "База лидов"

    # ---- title (A1:LAST_COL 1) ----
    ws.merge_cells(f"A1:{LAST_COL}1")
    t = ws["A1"]
    t.value = f"База лидов — {industry} (ИИ для бизнеса)"
    t.font = Font(name="Arial", size=13, bold=True, color=WHITE)
    t.fill = TITLE_FILL
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25.5

    # ---- header row 2 ----
    for i, h in enumerate(HEADERS):
        c = ws.cell(row=2, column=i + 1, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[2].height = 31.5

    # ---- column widths ----
    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w

    # ---- data rows (from row 3) ----
    wrap_cols = {3, 9, 10, 12}  # C, I, J, L — длинный текст
    for idx, ld in enumerate(leads, start=1):
        r = idx + 2
        row_vals = [
            idx,
            ld.get("name", ""),
            ld.get("niche", ""),
            ld.get("website", ""),
            ld.get("phone", ""),
            ld.get("email", ""),
            ld.get("contact_person", ""),
            ld.get("source", ""),
            ld.get("pain", ""),
            ld.get("offer", ""),
            ld.get("status", ""),
            ld.get("next_step", ""),
            ld.get("_lpr_post", ""),
            ld.get("_inn", ""),
            ld.get("_region", ""),
            (ld.get("_revenue") if isinstance(ld.get("_revenue"), (int, float)) else ""),
            ld.get("_revenue_source_url", ""),
            _confidence_label(ld),
        ]
        for ci, val in enumerate(row_vals, start=1):
            c = ws.cell(row=r, column=ci, value=val)
            c.font = Font(name="Arial", size=10)
            c.border = BORDER
            wrap = ci in wrap_cols
            horiz = "center" if ci in (1, 5) else "left"
            c.alignment = Alignment(horizontal=horiz, vertical="top", wrap_text=wrap)
            if idx % 2 == 0:
                c.fill = ZEBRA
            # колонка F (Email): пометить тип почты цветом + комментарием
            if ci == 6 and val and ld.get("_email_kind"):
                _mark_email(c, ld)
            # колонка P (Выручка): числовой формат + год/источник в комментарии
            if ci == 16 and isinstance(val, (int, float)) and val != "":
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right", vertical="top")
                yr = ld.get("_revenue_year") or "?"
                src = ld.get("_revenue_src")
                if src == "girbo":
                    note = (f"Выручка за {yr} (стр. 2110)\n"
                            "Источник: ГИР БО ФНС (bo.nalog.gov.ru)")
                elif src == "dadata":
                    note = f"Доходы за {yr} (данные ФНС, Dadata)"
                else:
                    note = f"Выручка за {yr}"
                c.comment = Comment(note, "lead-finder")
            elif ci == 16 and ld.get("_revenue_note"):
                c.comment = Comment(ld["_revenue_note"], "lead-finder")
            # колонка Q (Источник выручки): кликабельная ссылка + ГОД выручки
            if ci == 17 and val:
                c.hyperlink = val
                src_name = ld.get("_revenue_source_name") or "источник"
                yr = ld.get("_revenue_year")
                c.value = f"{src_name}, {yr}" if yr else src_name
                c.font = Font(name="Arial", size=10, color="FF1F6FB2", underline="single")
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 30

    ws.freeze_panes = "A3"
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    import json, sys
    leads = json.load(open(sys.argv[1], encoding="utf-8"))
    out = sys.argv[2] if len(sys.argv) > 2 else "test_leads.xlsx"
    industry = sys.argv[3] if len(sys.argv) > 3 else "Тест"
    print("written:", build(leads, industry, out))
