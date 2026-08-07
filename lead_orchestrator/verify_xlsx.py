# -*- coding: utf-8 -*-
r"""
verify_xlsx — добор неподтверждённых адресов из готовой выгрузки .xlsx через
облачный верификатор (email_guess._mev_fill / MyEmailVerifier).

Зачем отдельный вход. Выгрузки вида «Почты руководителей филиалов и топ-менеджмента»
уже собраны и разосланы по рукам: гонять ради них весь пайплайн заново незачем, а
подтвердить оставшиеся адреса надо. Зелёная заливка колонки «Почта» = адрес уже
подтверждён почтовым сервером; трогаем ТОЛЬКО НЕзелёные строки.

⚠️ Комплаенс. Тот же гейт, что и в пайплайне (`email_guess._mev_allowed`), но здесь
он опаснее: в .xlsx нет колонки отрасли, а fail-closed без отрасли заблокировал бы
всё. Поэтому отрасль задаётся явно (`--industry`) и всё равно проверяется по
названию компании — на реальной выгрузке под меткой «строительство» лежали ФНПЦ
«Титан-Баррикады», ГосНИИ «Кристалл» и ОКБ «Новатор». Исключённые компании
печатаются поимённо: молча пропустить их нельзя, это и есть суть проверки.

Дефолт — СУХОЙ прогон: показать, что ушло бы наружу, и ничего не отправлять.
Реальный прогон — только явным `--run`.

  py verify_xlsx.py "D:\...\Почты руководителей.xlsx" --industry construction
  py verify_xlsx.py "...xlsx" --industry construction --run --limit 100
  py verify_xlsx.py "...xlsx" --industry construction --run --out "...проверено.xlsx"
"""
from __future__ import annotations

import argparse
import os
import sys
from collections import Counter

SCRIPTS = os.path.dirname(os.path.abspath(__file__))
if SCRIPTS not in sys.path:
    sys.path.insert(0, SCRIPTS)

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import project_env                              # noqa: F401 — тихо подгружает env/.env
except Exception:
    pass

import openpyxl                                     # noqa: E402
from openpyxl.styles import PatternFill             # noqa: E402
from openpyxl.utils import get_column_letter        # noqa: E402

import email_guess as EG                            # noqa: E402

# Заливка «подтверждён сервером» в самой выгрузке — стандартный зелёный Excel.
# openpyxl отдаёт rgb то с альфой (FFC6EFCE), то с нулями (00C6EFCE), а шестизначный
# цвет при записи сам дополняет до "00…" — сравниваем по последним шести знакам.
GREEN = "C6EFCE"
UNKNOWN_RGB = "FFEB9C"                              # жёлтый: спросили, ясности нет
# Подпись и заливка вердикта одной таблицей. «ok» красится ТЕМ ЖЕ зелёным, который
# читается на входе: размеченный сегодня адрес следующий прогон пропустит и второй
# раз за него не заплатит.
VERDICT_MARK = {
    "ok": ("подтверждён", GREEN),
    "no": ("НЕ отправлять — ящика нет", "FFC7CE"),
    "unknown": ("не подтверждён", UNKNOWN_RGB),
}


def log(message):
    print(message, flush=True)


def _col_index(ws, title):
    """Номер колонки по заголовку первой строки (1-based) или 0."""
    for cell in ws[1]:
        if str(cell.value or "").strip().lower() == title.lower():
            return cell.column
    return 0


def _is_green(cell):
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return False
    rgb = getattr(fill.fgColor, "rgb", None)
    return isinstance(rgb, str) and rgb.upper().endswith(GREEN)


def _text(row, col):
    """Ячейка строки как обрезанный текст; нет такой колонки в выгрузке — пусто."""
    return str(row[col - 1].value or "").strip() if col else ""


def read_rows(path, sheet=None):
    """Незелёные строки с адресом -> (книга, лист, список dict'ов по строке таблицы)."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    # Резолвим ровно те колонки, которые используем: остальное в выгрузке справочное
    cols = {name: _col_index(ws, name) for name in ("Компания", "Почта", "Домен")}
    missing = [name for name in ("Компания", "Почта") if not cols[name]]
    if missing:
        raise SystemExit(f"в листе «{ws.title}» нет колонок: {', '.join(missing)}")

    rows = []
    for row in ws.iter_rows(min_row=2):
        mail_cell = row[cols["Почта"] - 1]
        email = str(mail_cell.value or "").strip().lower()
        if "@" not in email:
            continue                               # «нет кандидатов» — проверять нечего
        if _is_green(mail_cell):
            continue                               # уже подтверждён сервером
        rows.append({
            "row": mail_cell.row,
            "company": _text(row, cols["Компания"]),
            "email": email,
            # колонки «Домен» может не быть, а ячейка бывает пустой — тогда берём
            # домен из самого адреса: блок-лист доменов обязан работать всегда
            "domain": _text(row, cols["Домен"]).lower() or email.split("@")[-1],
        })
    return wb, ws, rows


def _gate_lead(company, industry):
    """Лид для гейта email_guess из того, что есть в таблице.

    Из трёх гейтов _mev_allowed здесь работают два: отрасль задаёт оператор, а ОКВЭД
    в выгрузке нет вовсе — и это ещё одна причина не ослаблять словарный. Название
    кладём и в name, и в niche: _mev_allowed сканирует оба поля, и лишняя копия
    дешевле пропущенного ФНПЦ."""
    return {"name": company, "_industry": industry, "niche": company,
            "_okved_descr": ""}


def split_by_gate(rows, industry):
    """Разделить строки на выпускаемые наружу и отсечённые гейтом.

    Отсечённые — счётчик {(компания, причина): адресов}: наружу от них нужно только
    имя с причиной, их печатает main поимённо."""
    allowed, blocked = [], Counter()
    for item in rows:
        ok, why = EG._mev_allowed(_gate_lead(item["company"], industry), item["domain"])
        if ok:
            allowed.append(item)
        else:
            blocked[(item["company"], why)] += 1
    return allowed, blocked


def unique_emails(allowed, limit=0):
    """Адреса, которые реально уйдут наружу: по одному вызову на адрес, порядок
    стабильный. Один и тот же список считает сухой прогон и боевой — иначе оператору
    показали бы одно, а отправили другое."""
    todo = sorted({item["email"] for item in allowed})
    return todo[:limit] if limit else todo


def verify(allowed, industry, limit=0):
    """Прогнать адреса через облачный слой. Возвращает {адрес: {verdict, note}}.

    Группируем по домену ради читаемого лога; паузы под лимит 30 запросов/мин,
    суточную квоту и кэш держит сам _mev_fill."""
    company_of = {item["email"]: item["company"] for item in allowed}
    by_domain = {}
    for email in unique_emails(allowed, limit):
        by_domain.setdefault(email.split("@")[-1], []).append(email)

    out = {}
    for domain, addrs in sorted(by_domain.items()):
        # Форма out из verify_addresses: _mev_fill читает domain/catch_all/verdicts,
        # остальное дописывает сам. Лид отдаём снова — по нему _mev_fill спросит
        # гейт ещё раз, уже перед самой сетью.
        state = {"domain": domain, "probed": False, "catch_all": None, "reason": "",
                 "verdicts": {a: {"verdict": "unknown", "trusted": False, "note": ""}
                              for a in addrs}}
        EG._mev_fill(state, _gate_lead(company_of[addrs[0]], industry), log=None)
        got = {a: {"verdict": row["verdict"], "note": row["note"]}
               for a, row in state["verdicts"].items()}
        out.update(got)
        tally = Counter(row["verdict"] for row in got.values())
        log(f"  {domain}: {len(addrs)} адр. -> "
            + ", ".join(f"{kind}×{n}" for kind, n in tally.most_common()))
    return out


def write_back(wb, ws, rows, verdicts, out_path):
    """Проставить вердикты в копию книги: цвет заливки + колонки вердикта и ноты."""
    col_verdict = _col_index(ws, "Проверка MyEmailVerifier") or ws.max_column + 1
    col_note = col_verdict + 1
    col_mail = _col_index(ws, "Почта")
    ws.cell(row=1, column=col_verdict, value="Проверка MyEmailVerifier")
    ws.cell(row=1, column=col_note, value="Пояснение MyEmailVerifier")
    ws.column_dimensions[get_column_letter(col_verdict)].width = 26
    ws.column_dimensions[get_column_letter(col_note)].width = 60

    changed = 0
    for item in rows:
        got = verdicts.get(item["email"])
        if not got:
            continue
        label, rgb = VERDICT_MARK.get(got["verdict"], (got["verdict"], UNKNOWN_RGB))
        fill = PatternFill("solid", fgColor=rgb)
        cell = ws.cell(row=item["row"], column=col_verdict, value=label)
        cell.fill = fill
        ws.cell(row=item["row"], column=col_note, value=got["note"])
        if got["verdict"] != "unknown":
            # Красим и саму почту: именно по её цвету файл читают глазами
            ws.cell(row=item["row"], column=col_mail).fill = fill
            changed += 1
    wb.save(out_path)
    return changed


def main():
    ap = argparse.ArgumentParser(
        description="Добор неподтверждённых адресов из .xlsx через MyEmailVerifier")
    ap.add_argument("xlsx", help="файл выгрузки")
    ap.add_argument("--sheet", help="лист (по умолчанию первый)")
    ap.add_argument("--industry", required=True,
                    help="отрасль прогона (ключ source_rusprofile.INDUSTRY): "
                         "в таблице её нет, а без неё гейт запрещает всё")
    ap.add_argument("--run", action="store_true",
                    help="реально обращаться к сервису (без флага — сухой прогон)")
    ap.add_argument("--limit", type=int, default=0, help="максимум адресов за прогон")
    ap.add_argument("--out", help="куда сохранить размеченную копию (по умолчанию "
                                  "рядом, с суффиксом «_проверено»)")
    a = ap.parse_args()

    wb, ws, rows = read_rows(a.xlsx, a.sheet)
    log(f"[файл] {a.xlsx} — лист «{ws.title}»")
    log(f"[строки] незелёных с адресом: {len(rows)}; "
        f"уникальных адресов: {len({item['email'] for item in rows})}")

    allowed, blocked = split_by_gate(rows, a.industry)

    if blocked:
        log(f"\n[гейт] НЕ выпускаются наружу ({sum(blocked.values())} строк):")
        for (company, why), count in sorted(blocked.items()):
            log(f"  • {company} — {why} ({count} адр.)")
    uniq = unique_emails(allowed, a.limit)
    log(f"\n[к проверке] компаний: {len({item['company'] for item in allowed})}, "
        f"уникальных адресов: {len(uniq)}"
        + (f" (ограничено --limit {a.limit})" if a.limit else ""))

    if not a.run:
        log("\nСУХОЙ ПРОГОН — наружу не ушло ничего. Адреса, которые ушли бы:")
        for email in uniq:
            log(f"  {email}")
        log("\nРеальный прогон: добавить --run (нужен EMAIL_MEV_API_KEY и "
            "EMAIL_MEV_ENABLE=1)")
        return 0

    if not EG.mev_enabled():
        raise SystemExit(
            "[стоп] облачный слой выключен. Нужны EMAIL_MEV_ENABLE=1 и "
            "EMAIL_MEV_API_KEY (ключ — в env/.env, в git он не попадает).")
    left, why = EG.mev_credits()
    if left is None:
        raise SystemExit(f"[стоп] баланс не получен: {why}")
    log(f"\n[кредиты] на счету {left}, суточный лимит "
        f"{EG.mev_policy()['daily_limit']}")
    if left < len(uniq):
        log(f"⚠️ кредитов меньше, чем адресов ({left} < {len(uniq)}) — "
            f"остальные останутся непроверенными")

    log("\n[прогон]")
    verdicts = verify(allowed, a.industry, a.limit)

    out_path = a.out or (os.path.splitext(a.xlsx)[0] + "_проверено.xlsx")
    changed = write_back(wb, ws, rows, verdicts, out_path)
    got = [v["verdict"] for v in verdicts.values()]
    log(f"\n[итог] спрошено {len(verdicts)}; подтверждено {got.count('ok')}, "
        f"«ящика нет» {got.count('no')}, осталось неясными {got.count('unknown')}")
    log(f"[файл] размечено строк {changed} -> {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
