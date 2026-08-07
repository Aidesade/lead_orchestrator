# -*- coding: utf-8 -*-
r"""
Офлайн-тест verify_xlsx — без сети и без обращений к облачному сервису.

Держит три вещи, которые в этом модуле только и могут сломаться незаметно:
  * «зелёный» распознаётся по заливке (openpyxl отдаёт rgb то как FFC6EFCE, то
    как 00C6EFCE) — ошибка здесь молча погнала бы уже подтверждённые адреса на
    повторную платную проверку;
  * комплаенс-гейт применяется к КАЖДОЙ строке по названию компании. В таблице нет
    колонки отрасли, поэтому словарный гейт — единственное, что отделяет ФНПЦ и
    ОКБ от дорожных ДРСУ. Проверяем на реальных названиях из выгрузки;
  * круг замкнут: чем write_back красит подтверждённый адрес, то read_rows на
    следующем прогоне и считает зелёным.

Запуск: py test_verify_xlsx.py
"""
import os
import pathlib
import sys
import tempfile

HERE = pathlib.Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import openpyxl
from openpyxl.styles import PatternFill

import verify_xlsx as VX

_fails = []


def check(name, cond, detail=""):
    if cond:
        print(f"[OK ] {name}")
    else:
        print(f"[FAIL] {name}" + (f" — {detail}" if detail else ""))
        _fails.append(name)


HEADERS = ["Компания", "ИНН", "ФИО", "Должность", "Уровень", "Почта",
           "Тип адреса", "Статус проверки", "Уверенность", "Откуда ФИО",
           "Страница-источник", "Домен"]
# Компании — из реальной выгрузки «Почты руководителей филиалов и топ-менеджмента»
ROWS = [
    ("АО Мостострой-11", "8617001665", "Руссу Николай Александрович",
     "tmn@ms11.ru", "ms11.ru", "FFC6EFCE"),          # зелёный, альфа FF
    ("АО РИК Автодор", "1419005577", "Иванов Иван Иванович",
     "aa@ricsakha.ru", "ricsakha.ru", "00C6EFCE"),   # зелёный, альфа 00
    ("АО Брянскавтодор", "3250510627", "Куст Андрей",
     "a.kust@avtodor32.ru", "avtodor32.ru", None),   # без заливки — к проверке
    ("АО ПО Возрождение", "7811062995", "Абрамчук Андрей",
     "abramchukan@vozr.ru", "vozr.ru", "00F2F2F2"),  # серый — тоже к проверке
    ("АО ФНПЦ Титан-Баррикады", "3442110950", "Петров Пётр",
     "petrov@titan-barrikady.ru", "titan-barrikady.ru", None),
    ("АО ОКБ Новатор", "6673092045", "Сидоров Сидор",
     "sidorov@novator.ru", "novator.ru", None),
    ("АО Госнии Кристалл", "5249116549", "Комаров А",
     "a.komarov@niikristall.ru", "niikristall.ru", None),
    ("ООО Красноармейское ДРСУ", "2336018775", "Ким Ким",
     "", "", None),                                  # нет адреса — строка не в счёт
]


def _make_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Руководители"
    ws.append(HEADERS)
    for company, inn, fio, mail, domain, rgb in ROWS:
        ws.append([company, inn, fio, "Директор", "центральный аппарат", mail,
                   "", "", "", "", "", domain])
        if rgb:
            ws.cell(row=ws.max_row, column=6).fill = PatternFill("solid", fgColor=rgb)
    wb.save(path)


def test_read_rows():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "leaders.xlsx")
        _make_book(path)
        _wb, ws, rows = VX.read_rows(path)

        check("лист найден", ws.title == "Руководители", ws.title)
        mail_col = VX._col_index(ws, "Почта")
        check("колонка «Почта» опознана по заголовку", mail_col == 6, str(mail_col))
        names = [r["company"] for r in rows]
        check("зелёные строки пропущены (обе формы rgb)",
              "АО Мостострой-11" not in names and "АО РИК Автодор" not in names,
              str(names))
        check("строка без адреса пропущена",
              "ООО Красноармейское ДРСУ" not in names, str(names))
        check("незелёные с адресом взяты", len(rows) == 5, f"{len(rows)}: {names}")
        check("серая заливка зелёной не считается",
              "АО ПО Возрождение" in names, str(names))
        check("адрес приведён к нижнему регистру",
              all(r["email"] == r["email"].lower() for r in rows))
        check("домен подставлен из колонки",
              rows[0]["domain"] == "avtodor32.ru", rows[0]["domain"])


def test_gate():
    """Главное: под меткой «строительство» едут ФНПЦ, ОКБ и ГосНИИ."""
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "leaders.xlsx")
        _make_book(path)
        _wb, _ws, rows = VX.read_rows(path)
        allowed, blocked = VX.split_by_gate(rows, "construction")

        out = sorted(item["company"] for item in allowed)
        stopped = sorted(company for company, _why in blocked)
        check("оборонные отсечены все три",
              stopped == ["АО Госнии Кристалл", "АО ОКБ Новатор",
                          "АО ФНПЦ Титан-Баррикады"], str(stopped))
        check("гражданские дорожники пропущены",
              out == ["АО Брянскавтодор", "АО ПО Возрождение"], str(out))
        check("причина отсечения названа",
              all(why for _company, why in blocked), str(list(blocked)))
        # Сухой прогон печатает ровно тот список, которым потом пойдёт боевой
        todo = VX.unique_emails(allowed)
        check("адреса дедуплицированы, порядок стабильный",
              todo == ["a.kust@avtodor32.ru", "abramchukan@vozr.ru"], str(todo))
        check("--limit режет тот же список",
              VX.unique_emails(allowed, 1) == todo[:1], str(todo[:1]))

        # Отрасль из блок-листа запрещает ВСЁ, даже гражданские названия
        allowed, blocked = VX.split_by_gate(rows, "opk")
        check("отрасль opk запрещает весь файл", allowed == [], str(allowed))

        # Отрасль не задана — fail-closed, наружу не уходит ничего
        allowed, _blocked = VX.split_by_gate(rows, "")
        check("без отрасли наружу не уходит ничего", allowed == [], str(allowed))


def test_write_back():
    """Замыкание круга: чем красит write_back, то и читает read_rows.

    Разойдись эти два цвета — следующий прогон снова заплатил бы за уже
    подтверждённые адреса, а вердикт «ящика нет» пришлось бы искать глазами."""
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "leaders.xlsx")
        marked = os.path.join(tmp, "leaders_проверено.xlsx")
        _make_book(path)
        wb, ws, rows = VX.read_rows(path)
        changed = VX.write_back(wb, ws, rows, {
            "a.kust@avtodor32.ru": {"verdict": "ok", "note": "ящик существует"},
            "abramchukan@vozr.ru": {"verdict": "unknown", "note": "greylisted"},
        }, marked)
        check("размечена только строка с окончательным вердиктом", changed == 1,
              str(changed))

        _wb, _ws, again = VX.read_rows(marked)
        left = [item["email"] for item in again]
        check("подтверждённый адрес в работу больше не берётся",
              "a.kust@avtodor32.ru" not in left, str(left))
        check("неясный адрес остаётся к проверке",
              "abramchukan@vozr.ru" in left, str(left))


def main():
    for test in (test_read_rows, test_gate, test_write_back):
        print(f"\n--- {test.__name__} ---")
        test()
    print()
    if _fails:
        print(f"test_verify_xlsx: ПРОВАЛЕНО {len(_fails)}: {', '.join(_fails)}")
        return 1
    print("test_verify_xlsx: OK — зелёные строки не перепроверяются, "
          "оборонные компании наружу не уходят, без отрасли гейт закрыт")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
